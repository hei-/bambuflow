"""Office 与表格文件的解析：Word（python-docx）、Excel（openpyxl）、CSV（标准库）。

这些格式没有固定的分页，块的位置只记字符区间（以解析后的全文为准）。
Word 按标题样式维护标题路径，表格按行保留并以第一行为表头；Excel 每个工作表一个表格，
取单元格的计算值而不是公式（技术设计文档 5.5），百分比按单元格格式还原。
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile

import docx
import openpyxl
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from bambuflow_processing.blocks import Block, Outline, ParseError, join_lines, table_text

# 超过 10 万行的表只索引结构和摘要（技术设计文档 5.5），摘要依赖增强阶段，P1 先拒收
MAX_ROWS = 100_000
_HEADING_STYLE = re.compile(r"^Heading (\d)$")


def parse_docx(data: bytes, title: str) -> list[Block]:
    try:
        d = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise ParseError(f"Word 文件无法打开: {e}") from e
    outline = Outline()
    out: list[Block] = []
    for item in d.iter_inner_content():
        if isinstance(item, Paragraph):
            text = join_lines(item.text.splitlines())
            if not text:
                continue
            level = _heading_level(item)
            if level:
                out.append(Block("heading", text, outline.heading(level, text)))
            else:
                out.append(Block("paragraph", text, outline.path or (title,)))
        elif isinstance(item, Table):
            rows = table_text([cell.text for cell in row.cells] for row in item.rows)
            if len(rows) >= 2:
                out.append(Block("table", "\n".join(rows), outline.path or (title,), rows=rows))
    return out


def _heading_level(p: Paragraph) -> int:
    name = p.style.name if p.style is not None and p.style.name else ""
    if name == "Title":
        return 1
    m = _HEADING_STYLE.match(name)
    # 文档标题占第 1 级，“标题 1”从第 2 级开始
    return int(m.group(1)) + 1 if m else 0


def parse_xlsx(data: bytes, title: str) -> list[Block]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        raise ParseError(f"Excel 文件无法打开: {e}") from e
    out: list[Block] = []
    try:
        for ws in wb.worksheets:
            raw: list[list[str]] = []
            for row in ws.iter_rows():
                if len(raw) > MAX_ROWS:
                    raise ParseError(f"工作表 {ws.title} 超过 {MAX_ROWS} 行")
                # 只读模式下空单元格（EmptyCell）没有数字格式
                raw.append([_cell_text(c.value, getattr(c, "number_format", None)) for c in row])
            rows = table_text(raw)
            path = (title, ws.title)
            out.append(Block("heading", ws.title, path))
            if len(rows) >= 2:
                out.append(Block("table", "\n".join(rows), path, rows=rows))
    finally:
        wb.close()
    return out


def _cell_text(value: object, number_format: str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, int | float):
        fmt = number_format or ""
        if fmt.endswith("%"):
            m = re.search(r"\.(0+)%$", fmt)
            return f"{value * 100:.{len(m.group(1)) if m else 0}f}%"
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def parse_csv(data: bytes, title: str) -> list[Block]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise ParseError("CSV 不是 UTF-8 编码") from e
    rows = table_text(csv.reader(io.StringIO(text)))
    if len(rows) > MAX_ROWS:
        raise ParseError(f"CSV 超过 {MAX_ROWS} 行")
    if len(rows) < 2:
        return []
    return [Block("table", "\n".join(rows), (title,), rows=rows)]
