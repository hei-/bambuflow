"""生成共享盘样例源（P1-T2）：文本 PDF、扫描件 PDF、Word、Excel、邮件、图片，只用于开发与评测。

内容是虚构的公司内部制度与通知，主题与 Wiki 样例相关但不重复。每个文件旁挂一个 <文件名>.meta.yaml，
模拟文档管理系统接口返回的 ID、属主、密级与 ACL（与 Wiki 的 front matter 字段相同）。
生成物不到 1 MB，随仓库提交（tools/test-upstream/fileshare/），
这样 chunk_id 与评测题不随库版本漂移；
改内容后重新运行：

    make run CMD='uv run --locked --all-packages python tools/test-upstream/gen_fileshare.py \
        tools/test-upstream/fileshare' OUT='tools/test-upstream/fileshare'

文档内容定义在 DOCS 里，e2e 按文件路径加载本模块读取清单与期望内容，所以模块顶层只导入标准库，
PyMuPDF、python-docx、openpyxl 在生成函数里导入。
"""

from __future__ import annotations

import csv
import io
import sys
import zipfile
from dataclasses import dataclass, field
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    import pymupdf

# 固定的时间，生成结果可重复
FIXED = (2026, 9, 1, 0, 0, 0)


@dataclass(frozen=True)
class Heading:
    text: str
    level: int = 2


@dataclass(frozen=True)
class Para:
    text: str


@dataclass(frozen=True)
class Table:
    rows: tuple[tuple[str, ...], ...]  # 第一行为表头
    widths: tuple[float, ...] = ()  # PDF 里的列宽（点）


Item = Heading | Para | Table


@dataclass(frozen=True)
class Doc:
    file: str
    id: str
    title: str
    owner: str
    classification: str
    acl: tuple[str, ...]
    updated_at: str
    kind: str  # pdf | scan | docx | xlsx | eml | png
    pages: tuple[tuple[Item, ...], ...] = ()  # pdf、scan：每页的内容；docx、png：只有一页
    header: str = ""  # PDF 页眉（去重检查用）
    sheets: tuple[tuple[str, Table], ...] = ()  # xlsx
    mail: dict[str, str] = field(default_factory=dict)  # eml 的信头与正文、附件


RATE_TABLE = Table(
    (
        ("产品", "期限", "利率下限", "定价审批"),
        ("个人消费贷款", "1 年以内（含）", "LPR + 0.35%", "支行"),
        ("个人消费贷款", "1 年以上至 5 年", "LPR + 0.55%", "分行信贷审批部"),
        ("个人住房贷款（首套）", "5 年以上", "LPR - 0.20%", "分行信贷审批部"),
        ("个人住房贷款（二套）", "5 年以上", "LPR + 0.30%", "分行信贷审批部"),
        ("个人经营贷款", "3 年以内（含）", "LPR + 0.80%", "分行信贷审批部"),
    ),
    (130, 110, 100, 111),
)

ASSET_TABLE = Table(
    (
        ("设备类别", "标准配置", "使用年限", "报废审批"),
        ("笔记本电脑", "14 英寸，16 GB 内存", "4 年", "部门负责人"),
        ("台式电脑", "27 英寸显示器，16 GB 内存", "5 年", "部门负责人"),
        ("打印复印一体机", "A3 幅面，网络打印", "6 年", "综合管理部"),
        ("工作手机", "公司统一采购机型", "3 年", "部门负责人"),
    )
)

DOCS = (
    Doc(
        file="credit-rate-notice-2026q4.pdf",
        id="fs-credit-rate-notice",
        title="关于 2026 年第四季度个人贷款利率定价的通知",
        owner="alice",
        classification="internal",
        acl=("group:credit-l2",),
        updated_at="2026-09-25T00:00:00Z",
        kind="pdf",
        header="信贷管理部发文 · 公司内部",
        pages=(
            (
                Heading("关于 2026 年第四季度个人贷款利率定价的通知", 1),
                Heading("一、适用范围"),
                Para(
                    "本通知适用于各分支机构 2026 年 10 月 1 日至 12 月 31 日新发放的个人消费贷款、"
                    "个人住房贷款和个人经营贷款。存量贷款的利率按原合同约定执行，不因本通知调整。"
                ),
                Heading("二、利率下限"),
                Para("各产品的执行利率不得低于下表所列下限，LPR 指同期限档次的贷款市场报价利率。"),
                RATE_TABLE,
            ),
            (
                Heading("三、优惠定价审批"),
                Para(
                    "确需低于利率下限的优惠定价，须由经办机构提交书面申请，报总部个人金融部审批。"
                    "单笔优惠幅度不超过 0.30 个百分点，同一客户在本季度内只能申请一次。"
                ),
                Heading("四、执行时间"),
                Para(
                    "本通知自 2026 年 10 月 1 日起执行，有效期至 2026 年 12 月 31 日。"
                    "此前下发的第三季度利率定价通知同时停止执行。"
                ),
            ),
        ),
    ),
    Doc(
        file="branch-fire-inspection-scan.pdf",
        id="fs-fire-inspection",
        title="关于开展 2026 年度营业网点消防安全检查的通知（扫描件）",
        owner="eve",
        classification="internal",
        acl=("group:all-staff",),
        updated_at="2026-09-18T00:00:00Z",
        kind="scan",
        header="运营管理部 · 公司内部",
        pages=(
            (
                Heading("关于开展 2026 年度营业网点消防安全检查的通知", 1),
                Heading("一、检查时间"),
                Para("2026 年 11 月 1 日至 11 月 20 日，各营业网点先自查，再由运营管理部抽查。"),
                Heading("二、检查内容"),
                Para(
                    "灭火器是否在有效期内，应急照明与疏散指示是否完好，"
                    "疏散通道是否畅通，监控录像是否保存不少于 90 天。"
                ),
            ),
            (
                Heading("三、整改要求"),
                Para(
                    "检查发现的问题须在 5 个工作日内整改完毕，"
                    "整改报告由网点负责人签字后报运营管理部安全保卫岗备案。"
                ),
                Heading("四、责任追究"),
                Para("逾期未整改或整改后复查仍不合格的网点，取消其本年度安全评优资格。"),
            ),
        ),
    ),
    Doc(
        file="it-asset-management.docx",
        id="fs-it-asset",
        title="办公设备与软件资产管理办法",
        owner="carl",
        classification="public",
        acl=("group:all-staff",),
        updated_at="2026-08-30T00:00:00Z",
        kind="docx",
        pages=(
            (
                Heading("办公设备与软件资产管理办法", 1),
                Heading("第一章 总则"),
                Para("为规范办公设备与软件资产的配发、使用和报废，保障信息安全，制定本办法。"),
                Heading("第二章 设备配发与报废"),
                Para("员工入职时按岗位配发办公设备，各类设备的标准配置、使用年限与报废审批如下："),
                ASSET_TABLE,
                Para("设备达到使用年限且无法修复的，由使用部门提出报废申请，按上表审批。"),
                Heading("第三章 软件管理"),
                Para(
                    "员工只能安装公司软件目录中的软件，严禁安装盗版软件。"
                    "需要目录外软件的，须提交申请，由信息科技部在 10 个工作日内完成安全评估。"
                ),
                Heading("第四章 遗失与赔偿"),
                Para(
                    "设备遗失须在 24 小时内报告信息科技部并申请远程锁定；"
                    "因个人原因遗失或损坏的，按设备账面净值赔偿。"
                ),
            ),
        ),
    ),
    Doc(
        file="risk-limit-monitor-2026q3.xlsx",
        id="fs-risk-limit-monitor",
        title="2026 年第三季度风险限额监控表",
        owner="bob",
        classification="confidential",
        acl=("group:risk-dept",),
        updated_at="2026-10-08T00:00:00Z",
        kind="xlsx",
        sheets=(
            (
                "限额监控",
                Table(
                    (
                        ("限额指标", "限额值", "9 月末实际值", "使用率", "状态"),
                        ("单一集团客户授信集中度", "15%", "11.2%", "0.747", "正常"),
                        ("前十大客户授信集中度", "50%", "46.5%", "0.93", "预警"),
                        ("房地产贷款占比", "25%", "21.2%", "0.848", "正常"),
                        (
                            "交易账户风险价值（99%，1 天）",
                            "8000 万元",
                            "6900 万元",
                            "0.863",
                            "正常",
                        ),
                    )
                ),
            ),
            (
                "预警处置",
                Table(
                    (
                        ("预警级别", "触发条件", "处置要求"),
                        ("黄色预警", "使用率达到 90%", "3 个工作日内提交压降方案"),
                        ("红色预警", "使用率达到 100%", "暂停新增业务并报风险管理委员会"),
                    )
                ),
            ),
        ),
    ),
    Doc(
        file="aml-training-notice.eml",
        id="fs-aml-training-mail",
        title="关于 2026 年反洗钱专项培训安排的通知",
        owner="bob",
        classification="internal",
        acl=("group:credit-l2", "group:risk-dept"),
        updated_at="2026-10-09T00:00:00Z",
        kind="eml",
        mail={
            "from": "合规部 <compliance@corp.example>",
            "to": "信贷二级群组 <credit-l2@corp.example>, 风险管理部 <risk-dept@corp.example>",
            "date": "2026-10-09T09:30:00+08:00",
            "body": (
                "各位同事：\n\n"
                "为落实年度反洗钱培训要求，合规部将组织 2026 年反洗钱专项培训，安排如下。\n\n"
                "培训对象为信贷与风险条线全体员工，培训采用线上直播加线下考试的形式，"
                "具体场次见附件。\n\n"
                "考试合格线为 80 分，未通过的员工须在 30 日内参加补考。"
                "截至 12 月 31 日仍未完成培训的员工，暂停其可疑交易复核权限。\n\n"
                "请各部门于 10 月 20 日前报送参训名单。\n\n"
                "-- \n"
                "合规部 培训岗\n"
                "分机 8123\n\n"
                "免责声明：本邮件仅供指定收件人使用，如误收请立即删除。\n\n"
                "-----原始邮件-----\n"
                "发件人: 合规部\n"
                "主题: 2025 年反洗钱培训安排\n"
                "> 考试合格线为 70 分。\n"
            ),
            "attachment": "培训场次安排.csv",
        },
    ),
    Doc(
        file="visitor-notice.png",
        id="fs-visitor-notice",
        title="访客登记须知",
        owner="eve",
        classification="public",
        acl=("group:all-staff",),
        updated_at="2026-07-15T00:00:00Z",
        kind="png",
        pages=(
            (
                Heading("访客登记须知", 1),
                Para("一、访客须在前台出示有效身份证件并登记。"),
                Para("二、访客在办公区内须由接待员工全程陪同。"),
                Para("三、访客证当日有效，离开时交回前台。"),
                Para("四、访客不得进入机房和金库区域。"),
            ),
        ),
    ),
)

TRAINING_SCHEDULE = (
    ("场次", "日期", "时间", "形式"),
    ("第一场", "2026-11-05", "14:00–16:00", "线上直播"),
    ("第二场", "2026-11-12", "14:00–16:00", "线上直播"),
    ("补考", "2026-12-10", "09:30–11:00", "线下考场"),
)

# ---- PDF 与图片：PyMuPDF 自带的中文字体 china-s ----
# 字体嵌入页面并在保存前做子集化。直接用 fontname="china-s" 引用的是 PDF 标准中文字体，
# 不嵌入字形，阅读器按每个字符一个字宽排版，英文与数字的位置会和实际字形错开，
# 抽出的文字串到相邻单元格。
FONT = "cjk"

PAGE_W, PAGE_H, LEFT, RIGHT = 595.0, 842.0, 72.0, 523.0
SIZES = {1: 18.0, 2: 14.0}
BODY = 11.0


def _wrap(font: pymupdf.Font, text: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    cur = ""
    for ch in text:
        if font.text_length(cur + ch, fontsize=size) > width:
            lines.append(cur)
            cur = ch
        else:
            cur += ch
    return [*lines, cur] if cur else lines


def _draw_page(
    page: pymupdf.Page, items: tuple[Item, ...], header: str, no: int, total: int, top: float = 90
) -> None:
    import pymupdf

    font = pymupdf.Font("china-s")
    page.insert_font(fontname=FONT, fontbuffer=font.buffer)
    if header:
        page.insert_text((LEFT, 40), header, fontname=FONT, fontsize=9)
        page.insert_text(
            (LEFT, PAGE_H - 30), f"第 {no} 页 共 {total} 页", fontname=FONT, fontsize=9
        )
    y = top
    for it in items:
        if isinstance(it, Heading):
            size = SIZES[it.level]
            y += size * 0.6
            page.insert_text((LEFT, y + size), it.text, fontname=FONT, fontsize=size)
            y += size * 1.8
        elif isinstance(it, Para):
            for line in _wrap(font, it.text, BODY, RIGHT - LEFT):
                page.insert_text((LEFT, y + BODY), line, fontname=FONT, fontsize=BODY)
                y += BODY * 1.6
            y += BODY * 1.2
        else:
            widths = it.widths or tuple((RIGHT - LEFT) / len(it.rows[0]) for _ in it.rows[0])
            row_h = 24.0
            for r, row in enumerate(it.rows):
                x = LEFT
                for c, cell in enumerate(row):
                    rect = pymupdf.Rect(x, y, x + widths[c], y + row_h)
                    page.draw_rect(rect, color=(0, 0, 0), width=0.6)
                    size = 10.0
                    page.insert_text((x + 4, y + 16), cell, fontname=FONT, fontsize=size)
                    x += widths[c]
                y += row_h
                if r == len(it.rows) - 1:
                    y += 12
        if y > PAGE_H - 80:
            raise ValueError(f"第 {no} 页内容超出页面")


def _text_pdf(d: Doc) -> pymupdf.Document:
    import pymupdf

    doc = pymupdf.open()
    for no, items in enumerate(d.pages, start=1):
        page = doc.new_page(width=PAGE_W, height=PAGE_H)
        _draw_page(page, items, d.header, no, len(d.pages))
    return doc


def _save_pdf(doc: pymupdf.Document, path: Path, title: str) -> None:
    stamp = "D:20260901000000+00'00'"
    doc.subset_fonts()
    doc.set_metadata(
        {"title": title, "producer": "gen_fileshare", "creationDate": stamp, "modDate": stamp}
    )
    doc.save(path, garbage=3, deflate=True, no_new_id=True)


def write_pdf(d: Doc, path: Path) -> None:
    with _text_pdf(d) as doc:
        _save_pdf(doc, path, d.title)


def write_scan(d: Doc, path: Path) -> None:
    """扫描件：先排出文字版，再把每页渲染成灰度图，存成只有图片、没有文字层的 PDF。"""
    import pymupdf

    with _text_pdf(d) as src, pymupdf.open() as out:
        for page in src:
            pix = page.get_pixmap(dpi=150, colorspace=pymupdf.csGRAY)
            dst = out.new_page(width=page.rect.width, height=page.rect.height)
            dst.insert_image(dst.rect, stream=pix.tobytes("jpeg", jpg_quality=75))
        _save_pdf(out, path, d.title)


def write_png(d: Doc, path: Path) -> None:
    import pymupdf

    with pymupdf.open() as doc:
        page = doc.new_page(width=PAGE_W, height=300)
        _draw_page(page, d.pages[0], "", 1, 1, top=20)
        page.get_pixmap(dpi=144, colorspace=pymupdf.csGRAY).save(path)


# ---- Office ----


def _fixed_zip(path: Path) -> None:
    """Office 文件是 zip，库写入时带当前时间；统一改成固定时间，内容不变时字节也不变。"""
    with zipfile.ZipFile(path) as z:
        entries = [(info, z.read(info)) for info in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in entries:
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(fixed, data)


def write_docx(d: Doc, path: Path) -> None:
    import datetime as dt

    import docx

    doc = docx.Document()
    for it in d.pages[0]:
        if isinstance(it, Heading):
            doc.add_heading(it.text, level=0 if it.level == 1 else 1)
        elif isinstance(it, Para):
            doc.add_paragraph(it.text)
        else:
            table = doc.add_table(rows=len(it.rows), cols=len(it.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(it.rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    stamp = dt.datetime(*FIXED, tzinfo=dt.UTC)
    doc.core_properties.created = stamp
    doc.core_properties.modified = stamp
    doc.core_properties.title = d.title
    doc.save(str(path))
    _fixed_zip(path)


def write_xlsx(d: Doc, path: Path) -> None:
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    if wb.active is not None:
        wb.remove(wb.active)
    for name, table in d.sheets:
        ws = wb.create_sheet(name)
        for r, row in enumerate(table.rows, start=1):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c)
                # 使用率存成数值并设百分比格式，验证解析时按单元格格式还原成百分数
                if r > 1 and table.rows[0][c - 1] == "使用率":
                    cell.value = float(value)
                    cell.number_format = "0.0%"
                else:
                    cell.value = value
    stamp = dt.datetime(*FIXED)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.title = d.title
    wb.save(path)
    _fixed_zip(path)


# ---- 邮件 ----


def write_eml(d: Doc, path: Path) -> None:
    import datetime as dt

    m = d.mail
    msg = EmailMessage()
    msg["Subject"] = d.title
    msg["From"] = m["from"]
    msg["To"] = m["to"]
    msg["Date"] = format_datetime(dt.datetime.fromisoformat(m["date"]))
    msg["Message-ID"] = f"<{d.id}@corp.example>"
    msg.set_content(m["body"], charset="utf-8", cte="base64")
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerows(TRAINING_SCHEDULE)
    msg.add_attachment(
        buf.getvalue().encode(), maintype="text", subtype="csv", filename=m["attachment"]
    )
    msg.set_boundary("bambuflow-sample-boundary")
    path.write_bytes(msg.as_bytes())


WRITERS = {
    "pdf": write_pdf,
    "scan": write_scan,
    "docx": write_docx,
    "xlsx": write_xlsx,
    "eml": write_eml,
    "png": write_png,
}


def write_meta(d: Doc, path: Path) -> None:
    lines = [
        f"id: {d.id}",
        f"title: {d.title}",
        f"owner: {d.owner}",
        f"classification: {d.classification}",
        f"acl: [{', '.join(d.acl)}]",
        f"updated_at: {d.updated_at}",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    for old in out.iterdir():
        old.unlink()
    for d in DOCS:
        WRITERS[d.kind](d, out / d.file)
        write_meta(d, out / (d.file + ".meta.yaml"))


if __name__ == "__main__":
    main()
